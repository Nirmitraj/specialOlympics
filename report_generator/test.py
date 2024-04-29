from docx import Document

state='sc'

def fill_placeholders():
    # Open the existing Word document
    doc = Document('../static/Connecticut_Y14_State_Report_(21-22).docx')

    # Replace placeholders with live data
    live_data = {
        '##place##': 'Live data for placeholder 1',
        'placeholder2': 'Live data for placeholder 2',
        # Add more key-value pairs as needed
    }

    for paragraph in doc.paragraphs:
        for placeholder, data in live_data.items():
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, data)

    return doc

def write_word_doc(doc, filename):
    doc.save(filename)

# Fill placeholders in the template document
updated_doc = fill_placeholders()

# Write the updated Word document to a file
write_word_doc(updated_doc, '../static/updated_document.docx')
